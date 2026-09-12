"""PPTX Document Ingestion Module (§19)."""
from io import BytesIO
from typing import Any, Dict

from pptx import Presentation


def parse_pptx(file_bytes: bytes) -> Dict[str, Any]:
    prs = Presentation(BytesIO(file_bytes))
    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
        if texts:
            slides.append({"slide_number": i, "text": "\n".join(texts)})
    return {"status": "extracted", "slides": slides}
