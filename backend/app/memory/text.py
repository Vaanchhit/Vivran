"""File bytes -> plain lines, in reading order.

Separate from app/ingestion on purpose: that path chunks text for retrieval,
this one needs layout the chunker throws away — line breaks (a question starts
on its own line), DOCX tables (many papers put marks in a table column) and
PPTX slide boundaries and speaker notes.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from typing import List, Optional

SUPPORTED_FILE_TYPES = {"pdf", "docx", "pptx"}

# Below this many characters a PDF is almost certainly a scan or a photo of
# handwriting: there is no text layer to read.
MIN_TEXT_CHARS = 80


class UnreadableFile(ValueError):
    """The file could not be read as text. The message is shown to the teacher."""


@dataclass
class Slide:
    number: int
    title: str
    body: List[str]
    notes: str = ""


@dataclass
class ExtractedText:
    file_type: str
    lines: List[str]
    slides: List[Slide] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n".join(self.lines)


def file_type_for(filename: Optional[str]) -> Optional[str]:
    ext = (filename or "").rsplit(".", 1)[-1].lower() if "." in (filename or "") else ""
    return ext if ext in SUPPORTED_FILE_TYPES else None


def _clean(lines: List[str]) -> List[str]:
    return [" ".join(line.split()) for line in lines if line and line.strip()]


def _pdf(data: bytes) -> ExtractedText:
    from pypdf import PdfReader

    try:
        reader = PdfReader(BytesIO(data))
        lines: List[str] = []
        for page in reader.pages:
            lines.extend((page.extract_text() or "").splitlines())
    except Exception as e:  # pypdf raises a wide range of errors on damaged files
        raise UnreadableFile("This PDF could not be opened. It may be damaged or password-protected.") from e
    return ExtractedText("pdf", _clean(lines))


def _docx(data: bytes) -> ExtractedText:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        doc = Document(BytesIO(data))
    except Exception as e:
        raise UnreadableFile("This Word file could not be opened.") from e
    lines: List[str] = []
    # Body order, so a table of questions sits where it appears on the page.
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            lines.extend(Paragraph(child, doc).text.splitlines())
        elif tag == "tbl":
            for row in Table(child, doc).rows:
                cells = []
                for cell in row.cells:
                    text = " ".join(cell.text.split())
                    if text and (not cells or cells[-1] != text):  # merged cells repeat
                        cells.append(text)
                if cells:
                    lines.append("  ".join(cells))
    return ExtractedText("docx", _clean(lines))


def _pptx(data: bytes) -> ExtractedText:
    from pptx import Presentation

    try:
        prs = Presentation(BytesIO(data))
    except Exception as e:
        raise UnreadableFile("This PowerPoint file could not be opened.") from e
    slides: List[Slide] = []
    lines: List[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = " ".join(slide.shapes.title.text_frame.text.split())
        body: List[str] = []
        for shape in slide.shapes:
            if shape == slide.shapes.title or not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                text = " ".join("".join(run.text for run in para.runs).split())
                if text:
                    body.append(text)
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = " ".join(slide.notes_slide.notes_text_frame.text.split())
        slides.append(Slide(i, title, body, notes))
        lines.extend(_clean([title, *body]))
    return ExtractedText("pptx", lines, slides)


def extract(data: bytes, file_type: str) -> ExtractedText:
    if file_type == "pdf":
        out = _pdf(data)
    elif file_type == "docx":
        out = _docx(data)
    elif file_type == "pptx":
        out = _pptx(data)
    else:
        raise UnreadableFile("Upload a PDF, Word (.docx) or PowerPoint (.pptx) file.")
    if file_type == "pdf" and len(out.text) < MIN_TEXT_CHARS:
        raise UnreadableFile(
            "No text could be read from this file. Scanned pages and photos of handwriting "
            "can't be read yet; upload the typed original if you have it."
        )
    return out
